import fitz  # PyMuPDF : Rendu HD vectoriel du PDF
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image
from google import genai
import os
import sys
import re
import shutil
import time

def rogner_schema(image_path, coordonnees_norm):
    """Découpe l'image selon [ymin, xmin, ymax, xmax] (échelle 0-1000)."""
    try:
        img = Image.open(image_path)
        width, height = img.size
        ymin, xmin, ymax, xmax = coordonnees_norm

        # Conversion échelle 0-1000 vers pixels réels
        left = int((xmin / 1000.0) * width)
        top = int((ymin / 1000.0) * height)
        right = int((xmax / 1000.0) * width)
        bottom = int((ymax / 1000.0) * height)

        # Marge de sécurité (padding 3%) pour ne pas couper les bordures
        pad_w = int(width * 0.03)
        pad_h = int(height * 0.03)
        left = max(0, left - pad_w)
        top = max(0, top - pad_h)
        right = min(width, right + pad_w)
        bottom = min(height, bottom + pad_h)

        img_croppee = img.crop((left, top, right, bottom))
        path_crop = image_path.replace(".png", "_crop.png")
        img_croppee.save(path_crop)
        return path_crop
    except Exception as e:
        print(f"⚠️ Erreur de rognage : {e}")
        return image_path

# ==============================================================================
# 1. INITIALISATION & CONFIGURATION MULTI-MODÈLES
# ==============================================================================

API_KEY = os.environ.get("GEMINI_API_KEY")
if not API_KEY:
    print("❌ ERREUR CRITIQUE : La variable GEMINI_API_KEY n'est pas définie dans l'environnement.")
    print("👉 Exécutez dans votre terminal : export GEMINI_API_KEY='votre_clé_api'")
    sys.exit(1)

client = genai.Client(api_key=API_KEY)

# Liste prioritaire des modèles : bascule automatique en cas de 404 ou dépréciation
MODELS_PRIORITAIRES = [
    "gemini-3.5-flash-lite",
    "gemini-flash-lite-latest"
]

# ==============================================================================
# 2. MOTEUR STYLES & RENDU MS WORD ADVANCED
# ==============================================================================

def configurer_page(doc, top=0.8, bottom=0.8, left=0.8, right=0.8):
    """Configure des marges resserrées (0.8 pouce ≈ 2 cm) pour un rendu professionnel."""
    for section in doc.sections:
        section.top_margin = Inches(top)
        section.bottom_margin = Inches(bottom)
        section.left_margin = Inches(left)
        section.right_margin = Inches(right)


def parse_and_add_formatted_runs(paragraph, text):
    """
    Parseur syntaxique Markdown Inline :
    Transforme la syntaxe ***gras-italique***, **gras**, *italique* et `code`
    en véritables formats Word enrichis.
    """
    pattern = r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)'
    tokens = re.split(pattern, text)

    for token in tokens:
        if not token:
            continue
        if token.startswith('***') and token.endswith('***') and len(token) > 6:
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
        elif token.startswith('**') and token.endswith('**') and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            paragraph.add_run(token)


def ajouter_titre(doc, text, level, color_rgb):
    """Ajoute un titre stylisé avec palette de couleurs RGB et gestion du formatage inline."""
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    
    text_clean = re.sub(r'^#+\s*', '', text).strip()
    parse_and_add_formatted_runs(h, text_clean)
    
    for run in h.runs:
        run.font.color.rgb = color_rgb


def ajouter_encadre_callout(doc, text):
    """
    Génère un encadré de citation / définition à fond gris-bleu (#F0F4F8)
    avec une bordure latérale gauche bleu marine (#003366).
    """
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)

    cell = table.cell(0, 0)

    # Shading Fond
    shading_xml = parse_xml(r'<w:shd {} w:fill="F0F4F8"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading_xml)

    # Bordure Gauche
    borders_xml = parse_xml(
        r'<w:tcBorders {}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="003366"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'.format(nsdecls('w'))
    )
    cell._tc.get_or_add_tcPr().append(borders_xml)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)

    parse_and_add_formatted_runs(p, text)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)


# ==============================================================================
# 3. CONSIGNES IA PROFESSEUR & DÉTECTION SCHÉMA
# ==============================================================================

PROMPT_EXPERT = """
Tu es un ingénieur pédagogique d'élite, professeur agrégé et expert reconnu dans la numérisation, la restructuration et l'optimisation didactique de supports de cours universitaires.
Ton rôle est d'analyser l'image de cette diapositive de cours et de la convertir en un chapitre de cours rédigé, parfaitement structuré, hautement synthétique, professionnel et prêt à être converti en un document Microsoft Word (.docx).

======================================================================
1. RÈGLES D'ÉLIMINATION STRICTE DU BRUIT ET DES MÉTADONNÉES
======================================================================
- SUPPRESSION DU DÉCOR INSTITUTIONNEL : Ignore et élimine TOUTES les mentions récurrentes, filigranes, logos, en-têtes, pieds de page et masques de diapositives :
  * Noms d'établissements et universités (ex. "IUT", "Aix-Marseille Université", "Institut Universitaire de Technologie", etc.).
  * Intitulés de filières, cursus, diplômes ou modules (ex. "BUT MLT", "Management de la Logistique et des Transports", "TRM - Transport Routier de Marchandises", "CM 02", etc.).
  * Métadonnées de pagination, numéros de diapositives, dates, auteurs ou sous-titres fixes de bas d'écran (ex. "Page X / Y", "Diapositive Z").
- ZÉRO VERBOSITÉ SYSTÈME : Ne rédige AUCUNE phrase d'introduction ("Voici le résumé...", "Dans cette diapositive..."), AUCUNE conclusion ("En espérant que cela vous aide..."), et AUCUN bloc de code Markdown global (pas de balises ```markdown ... ``` autour de ta réponse). Renvoie directement le texte structuré.
- AUCUN SÉPARATEUR PARASITE : Ne génère aucune ligne horizontale de séparation de type `---` dans le corps du texte.

======================================================================
2. DIAGNOSTIC ET DÉTECTION VISION (OBLIGATOIRE EN PREMIÈRE LIGNE)
======================================================================
Examine avec une extrême précision l'organisation visuelle de la diapositive :
- Si la diapositive contient un schéma synoptique, un organigramme, une carte géographique, une illustration technique, un graphique, une chronologie ou un tableau comparatif complexe dont la capture apporte une vraie valeur graphique :
  --> Indique OBLIGATOIREMENT sur la TOUTE PREMIÈRE LIGNE :
      ---VISUEL_DETECTE: OUI---
  --> Indique OBLIGATOIREMENT sur la DEUXIÈME LIGNE les coordonnées du cadre englobant le schéma (bounding box) à l'échelle 0-1000 :
      CADRE_VISUEL: [ymin, xmin, ymax, xmax]
      (Où ymin, xmin, ymax, xmax sont des nombres entiers entre 0 et 1000 délimitant précisément la zone du schéma/dessin à découper).

- Si la diapositive ne contient que du texte brut, des puces simples, des titres isolés ou des listes sans valeur graphique ajoutée :
  --> Indique OBLIGATOIREMENT sur la TOUTE PREMIÈRE LIGNE :
      ---VISUEL_DETECTE: NON---

======================================================================
3. STRATÉGIE RÉDACTIONNELLE ET ANATOMIE DU CONTENU
======================================================================
--- CAS 1 : SI VISUEL_DETECTE EST OUI ---
Puisque le schéma découpé (croppé) sera inséré automatiquement dans le document Word juste au-dessus du texte :
- NE RECOPIE PAS mot à mot le texte qui est déjà clairement lisible à l'intérieur des bulles, cases ou flèches du schéma.
- Rédige une ANALYSE PÉDAGOGIQUE ET SYNTHÉTIQUE sous le visuel :
  * Explique le fonctionnement global du schéma, la dynamique du flux et les interactions logiques entre les concepts.
  * Extrais les points clés, définitions fondamentales, formules et règles à retenir absolument pour les révisions.
  * Titre cette section : `#### Explication & Synthèse du schéma`.

--- CAS 2 : SI VISUEL_DETECTE EST NON ---
Puisqu'aucune image ne sera insérée dans le document :
- Restructure et rédige L'INTÉGRALITÉ du contenu pédagogique sous la forme d'un cours rédigé, fluide, clair et exhaustif.
- Ne saute AUCUN concept, AUCUNE formule de calcul, AUCUN exemple chiffré, AUCUNE exception ni AUCUNE règle juridique.
- Reformule le style "prise de notes concises" sous forme de phrases complètes, élégantes et agréables à lire.

======================================================================
4. FORMULES MATHÉMATIQUES ET UNITÉS (SANS SYNTAXE LATEX)
======================================================================
- INTERDICTION ABSOLUE DU LATEX : N'utilise JAMAIS de symboles `$`, `$$`, `\\text{}`, `\\times`, `\\frac` ou de balises LaTeX.
- ÉCRITURE UNICODE LISIBLE : Rédige toutes les formules, calculs, unités et équivalences en texte brut Unicode immédiatement lisible dans Word.
  * Équivalences d'unités : `1 m³ = 330 kg`, `1 ml = 1 750 kg`, `527 € HT`.
  * Formules : `Poids Volumétrique (PV) = Volume (m³) × 330 kg`.
  * Opérateurs : Utilise toujours le symbole de multiplication `×` (et jamais `\\times` ni la lettre `x`).

======================================================================
5. NORMES DE FORMATAGE MARKDOWN STRICTES (POUR INTÉGRATION DOCX)
======================================================================
- HIÉRARCHIE DES TITRES :
  * Utilise `### ` pour le sujet principal ou concept clé de la diapositive.
  * Utilise `#### ` pour les sous-sections, étapes ou points spécifiques.
- MOTS CLÉS EN GRAS : Applique impérativement du gras (`**notion clé**` ou `**valeur**`) sur tous les termes techniques, définitions, chiffres clés, unités et règles.
- LISTES STRUCTUREES :
  * Utilise des puces simples (`- `) pour énumérer des caractéristiques, critères ou composantes.
  * Utilise des listes numérotées (`1. `, `2. `) pour détailler des processus chronologiques, algorithmes ou étapes méthodologiques.
- ENCADRÉS DE MISE EN VALEUR : Utilise la syntaxe de citation (`> `) pour isoler les définitions officielles, règles d'or, formules clés et avertissements "À RETENIR".

Applique rigoureusement ces consignes dès maintenant pour traiter la diapositive fournie.
"""

# ==============================================================================
# 4. GESTIONNAIRE DE REQUÊTES ROBUSTE (BACKOFF EXPONENTIEL & FALLBACK)
# ==============================================================================

def appeler_gemini_robuste(img_path, max_retries=5):
    """
    Appelle Gemini avec retry automatique (backoff exponentiel) en cas de surcharge (503/429)
    et bascule automatiquement de modèle si une erreur 404 ou d'incompatibilité survient.
    """
    img_pil = Image.open(img_path)

    for model_name in MODELS_PRIORITAIRES:
        delay = 2.0
        for attempt in range(max_retries):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=[img_pil, PROMPT_EXPERT]
                )
                if response and response.text:
                    return response.text.strip(), model_name
                raise ValueError("Réponse vide reçue de l'API Gemini.")

            except Exception as e:
                err_str = str(e)

                # Cas 1 : Modèle non trouvé (404) -> Essayer le modèle suivant dans la liste
                if "404" in err_str or "NOT_FOUND" in err_str:
                    print(f"   ⚠️ Modèle '{model_name}' indisponible (404). Bascule sur le modèle alternatif...")
                    break  # Sort de la boucle d'essais pour tester le modèle suivant

                # Cas 2 : Surcharge temporaire (503/429) -> Attente exponentielle (2s, 4s, 8s, 16s...)
                is_transient = any(code in err_str for code in ["503", "429", "UNAVAILABLE", "HIGH_DEMAND"])
                if is_transient and attempt < max_retries - 1:
                    print(f"   ⏳ Surcharge serveur (503/429). Pause de {delay:.1f}s (Tentative {attempt + 1}/{max_retries})...")
                    time.sleep(delay)
                    delay *= 2.0
                else:
                    # Si dernier essai ou erreur critique non gérée pour ce modèle
                    if model_name == MODELS_PRIORITAIRES[-1]:
                        raise e
                    break

    raise RuntimeError("Tous les modèles Gemini prioritaires ont échoué.")


# ==============================================================================
# 5. PIPELINE PRINCIPAL DE TRAITEMENT
# ==============================================================================

def traiter_pdf_ultime(pdf_path):
    if not os.path.exists(pdf_path):
        print(f"❌ Fichier introuvable : {pdf_path}")
        return

    nom_base = os.path.splitext(os.path.basename(pdf_path))[0]
    output_docx = f"{nom_base}_Cours_Ultime.docx"
    dossier_images = f"temp_img_{nom_base}"
    os.makedirs(dossier_images, exist_ok=True)

    print(f"\n============================================================")
    print(f"🚀 DÉMARRAGE DU TRAITEMENT : '{pdf_path}'")
    print(f"============================================================")

    pdf_doc = fitz.open(pdf_path)
    total_pages = len(pdf_doc)
    print(f"📊 Nombre total de pages / diapositives : {total_pages}\n")

    doc = Document()
    configurer_page(doc, top=0.8, bottom=0.8, left=0.8, right=0.8)

    # Entête de présentation
    titre_p = doc.add_heading(level=0)
    titre_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titre = titre_p.add_run(f"COURS COMPLET : {nom_base.replace('_', ' ').upper()}")
    run_titre.font.color.rgb = RGBColor(0, 51, 102)
    run_titre.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Retranscription intégrale enrichie par Vision AI & Intégration Chronologique des Schémas")
    r_sub.font.italic = True
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(110, 110, 110)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Palette graphique
    COLOR_H2 = RGBColor(0, 51, 102)    # Bleu Marine Deep
    COLOR_H3 = RGBColor(0, 102, 153)   # Bleu Océan
    COLOR_H4 = RGBColor(51, 102, 153)  # Bleu Acier

    for page_num in range(total_pages):
        print(f"📄 [{page_num + 1}/{total_pages}] Rendu HD & Analyse IA...")

        page = pdf_doc[page_num]
        
        # Rendu image HD (200 DPI = compromis parfait finesse vision / vitesse transfert)
        pix = page.get_pixmap(dpi=200)
        img_path = os.path.join(dossier_images, f"page_{page_num + 1}.png")
        pix.save(img_path)

        ajouter_titre(doc, f"Partie {page_num + 1}", level=2, color_rgb=COLOR_H2)

        texte_ia = ""
        try:
            texte_ia, modele_utilise = appeler_gemini_robuste(img_path)
        except Exception as e:
            print(f"⚠️ Mode de secours activé sur la page {page_num + 1} ({e})")
            # Fallback PyMuPDF direct : extraction brute si l'API est inaccessible
            texte_brut = page.get_text("text").strip()
            if texte_brut:
                p_fb = doc.add_paragraph()
                p_fb.add_run(texte_brut)

        if texte_ia:
            contient_visuel = "---VISUEL_DETECTE: OUI---" in texte_ia

            # 1. Extraction des coordonnées et rognage si un visuel est détecté
            path_a_inserer = img_path
            if contient_visuel:
                match_cadre = re.search(r'CADRE_VISUEL:\s*\[(\d+),\s*(\d+),\s*(\d+),\s*(\d+)\]', texte_ia)
                if match_cadre:
                    coords = [int(g) for g in match_cadre.groups()]
                    path_a_inserer = rogner_schema(img_path, coords)

            # 2. Nettoyage des balises système, des traits '---' et conversion LaTeX
            lignes_raw = texte_ia.split("\n")
            lignes_clean = []
            for l in lignes_raw:
                l_str = l.strip()
                # Filtrage des balises système et des séparateurs horizontaux
                if l_str.startswith("---VISUEL_DETECTE:") or l_str.startswith("VISUEL:") or l_str.startswith("CADRE_VISUEL:") or l_str == "---":
                    continue
                
                # Nettoyage de sécurité des balises LaTeX résiduelles -> Unicode
                l_clean = re.sub(r'\\text\{([^}]+)\}', r'\1', l)
                l_clean = re.sub(r'\$\$(.*?)\$\$', r'\1', l_clean)
                l_clean = re.sub(r'\$(.*?)\$', r'\1', l_clean)
                l_clean = l_clean.replace(r'\times', '×').replace(r'\euro', '€')
                lignes_clean.append(l_clean)

            texte_formate = "\n".join(lignes_clean).strip()

            # 3. SI UN SCHÉMA/DIAGRAMME EST DÉTECTÉ : Capture découpée insérée au centre
            if contient_visuel:
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)

                run_img = p_img.add_run()
                run_img.add_picture(path_a_inserer, width=Inches(5.6))

                p_leg = doc.add_paragraph()
                p_leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_leg.paragraph_format.space_after = Pt(10)
                r_leg = p_leg.add_run(f"Figure / Schéma synoptique — Support de cours (Page {page_num + 1})")
                r_leg.font.italic = True
                r_leg.font.size = Pt(8.5)
                r_leg.font.color.rgb = RGBColor(120, 120, 120)

            # 4. CONVERSION MARKDOWN -> ELEMENTS DOCX
            for ligne in texte_formate.split("\n"):
                l = ligne.strip()
                if not l:
                    continue

                if l.startswith("#### "):
                    ajouter_titre(doc, l.replace("#### ", ""), level=4, color_rgb=COLOR_H4)
                elif l.startswith("### "):
                    ajouter_titre(doc, l.replace("### ", ""), level=3, color_rgb=COLOR_H3)
                elif l.startswith("## "):
                    ajouter_titre(doc, l.replace("## ", ""), level=2, color_rgb=COLOR_H2)
                elif l.startswith("> "):
                    ajouter_encadre_callout(doc, l.replace("> ", ""))
                elif l.startswith("- ") or l.startswith("* "):
                    p = doc.add_paragraph(style='List Bullet')
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(3)
                    parse_and_add_formatted_runs(p, l[2:])
                elif re.match(r'^\d+\.\s', l):
                    p = doc.add_paragraph(style='List Number')
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(3)
                    texte_num = re.sub(r'^\d+\.\s', '', l)
                    parse_and_add_formatted_runs(p, texte_num)
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(4)
                    parse_and_add_formatted_runs(p, l)

        # Espacement entre les parties
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Tempo tactique de 3s entre chaque page
        time.sleep(3)

    # Enregistrement final
    doc.save(output_docx)

    # Nettoyage automatique des images temporaires
    try:
        shutil.rmtree(dossier_images)
    except Exception:
        pass

    print("============================================================")
    print(f"🎉 TRAITEMENT TERMINÉ AVEC SUCCÈS !")
    print(f"📄 Fichier Word généré : {output_docx}")
    print("============================================================\n")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        traiter_pdf_ultime(sys.argv[1])
    else:
        print("Usage : python3 traiter_pdf.py <nom_du_fichier.pdf>")
